import io
import base64
import logging
from typing import Dict, Any, List

logger = logging.getLogger(__name__)

def _period(item: Dict[str, Any]) -> str:
    if item.get("period"):
        return str(item["period"])
    start, end = item.get("start_date") or "", item.get("end_date") or ""
    if start and end:
        return f"{start} – {end}"
    return str(start or end or item.get("year") or "")

def _bullets(item: Dict[str, Any]) -> List[str]:
    bullets = item.get("bullets") or item.get("highlights")
    if isinstance(bullets, list) and bullets:
        return [str(b) for b in bullets if b]
    desc = item.get("description")
    if isinstance(desc, list):
        return [str(b) for b in desc if b]
    return [str(desc)] if desc else []

def _normalize(doc: Dict[str, Any]) -> Dict[str, Any]:
    """Accept both the flat resume-form payload and the ResumeArchitect
    snapshot ({"header": {...}, "professional_summary": ..., ...})."""
    header = doc.get("header") or {}
    personal = doc.get("personal") or {}
    professional = doc.get("professional") or {}
    skills = doc.get("skills") or []
    if isinstance(skills, dict):
        flat = []
        for v in skills.values():
            if isinstance(v, list):
                flat.extend(str(x) for x in v if x)
            elif v:
                flat.append(str(v))
        skills = flat
    return {
        "name": doc.get("name") or header.get("name") or personal.get("name") or personal.get("full_name") or "",
        "title": doc.get("title") or header.get("title") or professional.get("headline") or professional.get("title") or doc.get("target_role") or "",
        "email": doc.get("email") or header.get("email") or personal.get("email") or "",
        "phone": doc.get("phone") or header.get("phone") or personal.get("phone") or "",
        "location": doc.get("location") or header.get("location") or personal.get("location") or "",
        "nationality": doc.get("nationality") or "",
        "photo_url": doc.get("photo_url"),
        "summary": doc.get("summary") or doc.get("professional_summary") or professional.get("summary") or "",
        "experience": [
            {
                "role": e.get("role") or e.get("title") or "",
                "company": e.get("company") or "",
                "period": _period(e),
                "bullets": _bullets(e),
            }
            for e in (doc.get("experience") or []) if isinstance(e, dict)
        ],
        "education": [
            {
                "degree": e.get("degree") or e.get("field") or "",
                "institution": e.get("institution") or e.get("school") or "",
                "year": _period(e),
            }
            for e in (doc.get("education") or []) if isinstance(e, dict)
        ],
        "skills": [str(x) for x in skills if x] if isinstance(skills, list) else [],
        "languages": doc.get("languages") or [],
        "certifications": doc.get("certifications") or [],
    }

def _escape(value: Any) -> Any:
    """ReportLab Paragraph parses mini-markup; escape user text so '<' or '&' can't break rendering."""
    from xml.sax.saxutils import escape
    if isinstance(value, str):
        return escape(value)
    if isinstance(value, list):
        return [_escape(v) for v in value]
    if isinstance(value, dict):
        return {k: _escape(v) for k, v in value.items()}
    return value

class ResumeExporter:
    def __init__(self, renderer=None):
        self.renderer = renderer

    def export_pdf(self, doc: Dict[str, Any]) -> bytes:
        doc = _escape(_normalize(doc))
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle, Image as RLImage, KeepTogether
            )
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            from PIL import Image as PILImage

            buffer = io.BytesIO()
            # Standard A4: 595.27 x 841.89 pt. 36 pt margin leaves ~523 pt usable width.
            doc_template = SimpleDocTemplate(
                buffer,
                pagesize=A4,
                rightMargin=36,
                leftMargin=36,
                topMargin=36,
                bottomMargin=36
            )
            styles = getSampleStyleSheet()

            EU_BLUE = colors.HexColor('#0E4785')
            EU_LIGHT_BLUE = colors.HexColor('#EBF3FA')
            EU_BORDER = colors.HexColor('#C8D9EA')
            TEXT_DARK = colors.HexColor('#181816')
            TEXT_MUTED = colors.HexColor('#555550')

            name_style = ParagraphStyle(
                'EuropassName',
                parent=styles['Normal'],
                fontSize=20,
                leading=24,
                textColor=EU_BLUE,
                fontName='Helvetica-Bold'
            )
            badge_style = ParagraphStyle(
                'EuropassBadge',
                parent=styles['Normal'],
                fontSize=9,
                leading=11,
                textColor=colors.white,
                fontName='Helvetica-Bold'
            )
            section_hdr_style = ParagraphStyle(
                'EuropassSecHdr',
                parent=styles['Normal'],
                fontSize=11,
                leading=14,
                textColor=EU_BLUE,
                fontName='Helvetica-Bold',
                spaceBefore=10,
                spaceAfter=4
            )
            col_date_style = ParagraphStyle(
                'EuropassDate',
                parent=styles['Normal'],
                fontSize=8.5,
                leading=12,
                textColor=TEXT_MUTED,
                fontName='Helvetica-Bold'
            )
            col_title_style = ParagraphStyle(
                'EuropassRoleTitle',
                parent=styles['Normal'],
                fontSize=10,
                leading=13,
                textColor=EU_BLUE,
                fontName='Helvetica-Bold'
            )
            col_company_style = ParagraphStyle(
                'EuropassCompany',
                parent=styles['Normal'],
                fontSize=9,
                leading=12,
                textColor=TEXT_DARK,
                fontName='Helvetica-Bold'
            )
            body_style = ParagraphStyle(
                'EuropassBody',
                parent=styles['Normal'],
                fontSize=8.5,
                leading=12,
                textColor=TEXT_DARK,
                fontName='Helvetica'
            )
            contact_label_style = ParagraphStyle(
                'EuropassContactLabel',
                parent=styles['Normal'],
                fontSize=8,
                leading=11,
                textColor=EU_BLUE,
                fontName='Helvetica-Bold'
            )
            contact_val_style = ParagraphStyle(
                'EuropassContactVal',
                parent=styles['Normal'],
                fontSize=8.5,
                leading=11,
                textColor=TEXT_DARK,
                fontName='Helvetica'
            )

            story = []

            # 1. Header with Europass Badge and Personal Information
            name = doc["name"] or "Curriculum Vitae"
            headline = doc["title"]
            email = doc["email"]
            phone = doc["phone"]
            location = doc["location"]
            nationality = doc["nationality"]
            photo_url = doc.get("photo_url")

            # Check if photo is provided (data url base64)
            img_flowable = None
            if photo_url and photo_url.startswith("data:image"):
                try:
                    header, encoded = photo_url.split(",", 1)
                    image_data = base64.b64decode(encoded)
                    pil_img = PILImage.open(io.BytesIO(image_data))
                    img_buffer = io.BytesIO()
                    pil_img.save(img_buffer, format='PNG')
                    img_buffer.seek(0)
                    img_flowable = RLImage(img_buffer, width=65, height=80)
                except Exception:
                    img_flowable = None

            # Europass Top Branding Block
            eu_banner_data = [
                [
                    Paragraph("<b>europass</b> &nbsp;|&nbsp; CURRICULUM VITAE", badge_style),
                ]
            ]
            eu_banner_table = Table(eu_banner_data, colWidths=[523])
            eu_banner_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), EU_BLUE),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ]))
            story.append(eu_banner_table)
            story.append(Spacer(1, 8))

            # Header Details: Photo (if any) + Name + Contact block
            name_block = [
                Paragraph(name, name_style),
                Paragraph(headline, ParagraphStyle('Sub', parent=body_style, fontSize=9.5, fontName='Helvetica-Bold', textColor=TEXT_MUTED)),
                Spacer(1, 4),
                Table([
                    [Paragraph("<b>Email:</b>", contact_label_style), Paragraph(email, contact_val_style)],
                    [Paragraph("<b>Phone:</b>", contact_label_style), Paragraph(phone, contact_val_style)],
                    [Paragraph("<b>Address:</b>", contact_label_style), Paragraph(location, contact_val_style)],
                    [Paragraph("<b>Nationality:</b>", contact_label_style), Paragraph(nationality, contact_val_style)]
                ], colWidths=[65, 340 if not img_flowable else 275])
            ]

            if img_flowable:
                header_cols = [
                    [img_flowable],
                    name_block
                ]
                header_table = Table([header_cols], colWidths=[80, 443])
            else:
                header_table = Table([[name_block]], colWidths=[523])

            header_table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            story.append(header_table)
            story.append(Spacer(1, 10))
            story.append(HRFlowable(width="100%", thickness=1.5, color=EU_BLUE, spaceAfter=8))

            # 2. Personal Statement / Summary
            summary = doc.get("summary") or doc.get("professional", {}).get("summary")
            if summary:
                story.append(Paragraph("WORK OBJECTIVE & PROFESSIONAL SUMMARY", section_hdr_style))
                story.append(Paragraph(summary, body_style))
                story.append(Spacer(1, 8))
                story.append(HRFlowable(width="100%", thickness=0.5, color=EU_BORDER, spaceAfter=6))

            # 3. Work Experience (Official Europass 2-Column Standard)
            experiences = doc.get("experience") or []
            if experiences:
                story.append(Paragraph("WORK EXPERIENCE", section_hdr_style))
                exp_table_data = []
                for exp in experiences:
                    role = exp.get("role") or exp.get("title") or ""
                    comp = exp.get("company") or ""
                    period = exp.get("period") or ""
                    bullets = exp.get("bullets") or exp.get("highlights") or []

                    left_cell = [
                        Paragraph(period, col_date_style)
                    ]
                    right_cell = [
                        Paragraph(role, col_title_style),
                        Paragraph(comp, col_company_style),
                        Spacer(1, 2)
                    ]
                    for b in bullets:
                        right_cell.append(Paragraph(f"• {b}", body_style))
                    right_cell.append(Spacer(1, 6))

                    exp_table_data.append([left_cell, right_cell])

                exp_table = Table(exp_table_data, colWidths=[120, 403])
                exp_table.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 2),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (0, -1), 0),
                    ('RIGHTPADDING', (0, 0), (0, -1), 8),
                    ('LEFTPADDING', (1, 0), (1, -1), 8),
                    ('LINEBEFORE', (1, 0), (1, -1), 1, EU_BORDER),
                ]))
                story.append(exp_table)
                story.append(Spacer(1, 6))
                story.append(HRFlowable(width="100%", thickness=0.5, color=EU_BORDER, spaceAfter=6))

            # 4. Education & Training
            education = doc.get("education") or []
            if education:
                story.append(Paragraph("EDUCATION AND TRAINING", section_hdr_style))
                edu_table_data = []
                for edu in education:
                    deg = edu.get("degree") or edu.get("field") or ""
                    inst = edu.get("institution") or edu.get("school") or ""
                    yr = edu.get("year") or ""

                    left_cell = [
                        Paragraph(yr, col_date_style)
                    ]
                    right_cell = [
                        Paragraph(deg, col_title_style),
                        Paragraph(inst, col_company_style),
                        Spacer(1, 4)
                    ]
                    edu_table_data.append([left_cell, right_cell])

                edu_table = Table(edu_table_data, colWidths=[120, 403])
                edu_table.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 2),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (0, -1), 0),
                    ('RIGHTPADDING', (0, 0), (0, -1), 8),
                    ('LEFTPADDING', (1, 0), (1, -1), 8),
                    ('LINEBEFORE', (1, 0), (1, -1), 1, EU_BORDER),
                ]))
                story.append(edu_table)
                story.append(Spacer(1, 6))
                story.append(HRFlowable(width="100%", thickness=0.5, color=EU_BORDER, spaceAfter=6))

            # 5. Personal Skills & Competencies
            skills = doc.get("skills") or []
            if isinstance(skills, dict):
                skills = skills.get("technical") or []
            languages = doc.get("languages") or []

            skills_data = []

            # Mother tongue & languages
            if languages:
                skills_data.append([
                    Paragraph("<b>Languages</b>", col_date_style),
                    Paragraph(" • ".join(languages) if isinstance(languages, list) else str(languages), body_style)
                ])

            # Digital / Technical skills
            if skills:
                skills_data.append([
                    Paragraph("<b>Digital Competencies</b>", col_date_style),
                    Paragraph(" • ".join(skills), body_style)
                ])

            if skills_data:
                story.append(Paragraph("PERSONAL SKILLS & COMPETENCIES", section_hdr_style))
                skills_table = Table(skills_data, colWidths=[120, 403])
                skills_table.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ('LEFTPADDING', (0, 0), (0, -1), 0),
                    ('RIGHTPADDING', (0, 0), (0, -1), 8),
                    ('LEFTPADDING', (1, 0), (1, -1), 8),
                    ('LINEBEFORE', (1, 0), (1, -1), 1, EU_BORDER),
                ]))
                story.append(skills_table)

            # 6. Certifications & Accreditations
            certifications = doc.get("certifications") or []
            if isinstance(certifications, list) and certifications:
                story.append(Spacer(1, 6))
                story.append(HRFlowable(width="100%", thickness=0.5, color=EU_BORDER, spaceAfter=6))
                story.append(Paragraph("CERTIFICATIONS & ACCREDITATIONS", section_hdr_style))
                cert_data = []
                for cert in certifications:
                    if isinstance(cert, dict):
                        c_name = cert.get("name") or cert.get("title") or ""
                        c_issuer = cert.get("issuer") or cert.get("authority") or ""
                        c_date = cert.get("date") or cert.get("year") or ""
                        parts = [p for p in [c_name, c_issuer, str(c_date)] if p]
                        cert_text = " — ".join(parts) if parts else str(cert)
                    else:
                        cert_text = str(cert)
                    cert_data.append([
                        Paragraph("<b>Certification</b>", col_date_style),
                        Paragraph(f"• {cert_text}", body_style)
                    ])
                cert_table = Table(cert_data, colWidths=[120, 403])
                cert_table.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ('LEFTPADDING', (0, 0), (0, -1), 0),
                    ('RIGHTPADDING', (0, 0), (0, -1), 8),
                    ('LEFTPADDING', (1, 0), (1, -1), 8),
                    ('LINEBEFORE', (1, 0), (1, -1), 1, EU_BORDER),
                ]))
                story.append(cert_table)

            doc_template.build(story)
            return buffer.getvalue()

        except Exception:
            # No silent placeholder PDF: surface the failure so callers return an error.
            logger.exception("Resume PDF export failed")
            raise

    def export_docx(self, doc: Dict[str, Any]) -> bytes:
        doc = _normalize(doc)
        try:
            from docx import Document
            buffer = io.BytesIO()
            docx_doc = Document()
            docx_doc.add_heading(doc.get("name") or "Curriculum Vitae", 0)
            contact = " | ".join(x for x in [doc["title"], doc["email"], doc["phone"], doc["location"]] if x)
            if contact:
                docx_doc.add_paragraph(contact)
            summary = doc.get("summary")
            if summary:
                docx_doc.add_heading("Professional Summary", level=1)
                docx_doc.add_paragraph(summary)
            experiences = doc.get("experience") or []
            if experiences:
                docx_doc.add_heading("Work Experience", level=1)
                for exp in experiences:
                    p = docx_doc.add_paragraph()
                    heading = " — ".join(x for x in [exp["role"], exp["company"]] if x)
                    if exp["period"]:
                        heading += f" ({exp['period']})"
                    p.add_run(heading).bold = True
                    for b in exp["bullets"]:
                        docx_doc.add_paragraph(f"• {b}")
            education = doc.get("education") or []
            if education:
                docx_doc.add_heading("Education", level=1)
                for edu in education:
                    line = " — ".join(x for x in [edu["degree"], edu["institution"]] if x)
                    if edu["year"]:
                        line += f" ({edu['year']})"
                    docx_doc.add_paragraph(line)
            if doc["skills"]:
                docx_doc.add_heading("Skills", level=1)
                docx_doc.add_paragraph(" • ".join(doc["skills"]))
            certifications = doc.get("certifications") or []
            if certifications and isinstance(certifications, list):
                docx_doc.add_heading("Certifications & Accreditations", level=1)
                for cert in certifications:
                    if isinstance(cert, dict):
                        c_name = cert.get("name") or cert.get("title") or ""
                        c_issuer = cert.get("issuer") or cert.get("authority") or ""
                        c_date = cert.get("date") or cert.get("year") or ""
                        parts = [p for p in [c_name, c_issuer, str(c_date)] if p]
                        cert_text = " — ".join(parts) if parts else str(cert)
                    else:
                        cert_text = str(cert)
                    docx_doc.add_paragraph(f"• {cert_text}")
            docx_doc.save(buffer)
            return buffer.getvalue()
        except Exception:
            logger.exception("Resume DOCX export failed")
            raise

def export_resume_pdf(doc: Dict[str, Any]) -> bytes:
    return ResumeExporter().export_pdf(doc)

def export_resume_docx(doc: Dict[str, Any]) -> bytes:
    return ResumeExporter().export_docx(doc)
