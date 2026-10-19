import io
import logging
from typing import Dict, List, Any
import docx
from pdfminer.high_level import extract_text as extract_pdf_text
from ..ai_agents.profile_agent import ProfileAgent
from .agent_runs import run_warn, tool_call

logger = logging.getLogger(__name__)

MIN_TEXT_CHARS = 150  # below this, a "PDF" is treated as image-only (design tool / scan)
MAX_OCR_PAGES = 3


def _pdf_text_pdfminer(file_bytes: bytes) -> str:
    try:
        return (extract_pdf_text(io.BytesIO(file_bytes)) or "").strip()
    except Exception as e:
        logger.warning("pdfminer could not read the PDF: %s", type(e).__name__)
        return ""


def _pdf_text_pymupdf(file_bytes: bytes) -> str:
    try:
        import pymupdf  # PyMuPDF
    except ImportError:
        return ""
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
            return "\n".join(page.get_text("text", sort=True) for page in doc).strip()
    except Exception as e:
        logger.warning("PyMuPDF could not read the PDF: %s", type(e).__name__)
        return ""


def _pdf_page_images(file_bytes: bytes, max_pages: int = MAX_OCR_PAGES) -> List[str]:
    """Render the first pages as small base64 JPEGs (kept under ~170 KB for the vision API)."""
    import base64
    try:
        import pymupdf
    except ImportError:
        return []
    images = []
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
            for page in list(doc)[:max_pages]:
                for dpi, quality in ((110, 70), (90, 60), (72, 50)):
                    pix = page.get_pixmap(dpi=dpi)
                    data = pix.tobytes("jpeg", jpg_quality=quality)
                    if len(data) * 4 / 3 < 170_000:
                        break
                images.append(base64.b64encode(data).decode())
    except Exception as e:
        logger.warning("Could not render PDF pages: %s", type(e).__name__)
        return []
    return images


def _ocr_pdf_with_vision(file_bytes: bytes) -> str:
    from concurrent.futures import ThreadPoolExecutor
    import contextvars
    from .nvidia_llm import LLMError, call_nvidia_vision_ocr
    images = _pdf_page_images(file_bytes)
    if not images:
        return ""
    run_warn(f"The PDF has no readable text layer (design tool or scan); reading {len(images)} page image(s) with the NVIDIA vision model.")

    def read(i_img):
        i, img = i_img
        try:
            return call_nvidia_vision_ocr(img, page_label=f"page {i + 1}")
        except LLMError as e:
            run_warn(f"Vision model could not read page {i + 1}: {e}")
            return ""
    ctx = contextvars.copy_context()  # keep the dashboard run attached in worker threads
    with ThreadPoolExecutor(max_workers=len(images)) as pool:
        pages = list(pool.map(lambda x: ctx.copy().run(read, x), enumerate(images)))
    return "\n\n".join(p.strip() for p in pages if p and p.strip())


def _docx_text(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    for section in doc.sections:
        for p in list(section.header.paragraphs) + list(section.footer.paragraphs):
            if p.text.strip():
                parts.append(p.text)
    text = "\n".join(parts).strip()
    if len(text) < MIN_TEXT_CHARS:
        # text boxes / shapes (common in designed CV templates) are not in doc.paragraphs
        import re
        xml = doc.element.xml
        boxes = re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", xml)
        alt = " ".join(b for b in boxes if b.strip())
        if len(alt) > len(text):
            text = alt
    return text


def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> str:
    """Plain text from a PDF, DOCX or text file. Returns "" if nothing readable (never raw bytes)."""
    fn = filename.lower()
    if fn.endswith(".pdf") or file_bytes[:5] == b"%PDF-":
        text = _pdf_text_pdfminer(file_bytes)
        if len(text) < MIN_TEXT_CHARS:
            alt = _pdf_text_pymupdf(file_bytes)
            if len(alt) > len(text):
                text = alt
        if len(text) < MIN_TEXT_CHARS:
            ocr = _ocr_pdf_with_vision(file_bytes)
            if len(ocr) > len(text):
                text = ocr
        return text
    if fn.endswith(".docx"):
        try:
            return _docx_text(file_bytes)
        except Exception as e:
            logger.error("Could not read DOCX %s: %s", filename, type(e).__name__)
            return ""
    if fn.endswith(".doc"):
        run_warn("Old .doc files are not supported; save the CV as .docx or PDF.")
        return ""
    try:
        return file_bytes.decode("utf-8", errors="ignore").strip()
    except Exception:
        return ""


def extract_text_and_profile(file_bytes: bytes, filename: str):
    """Read the document once, then run the CV Agent. Returns (text, extraction)."""
    kind = (filename.rsplit(".", 1)[-1] if "." in filename else "text").lower()
    with tool_call(f"parser:{kind}", f"extract text from {filename}") as t:
        text = extract_text_from_file_bytes(file_bytes, filename)
        t["summary"] = f"{len(text)} characters of text extracted"
    if not text.strip():
        run_warn("No text could be extracted from the document (scanned image or empty file?)")
    return text, ProfileAgent().extract_from_text(text)


def extract_facts_from_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Extract full structured profile and facts from uploaded document bytes."""
    return extract_text_and_profile(file_bytes, filename)[1]

def confirm_facts(facts: List[Dict], approvals: Dict[str, bool]) -> List[Dict]:
    confirmed = []
    for f in facts:
        key = f.get("path") or f.get("fact_key")
        approved = approvals.get(key, True)
        f["verified"] = approved
        confirmed.append(f)
    return confirmed
